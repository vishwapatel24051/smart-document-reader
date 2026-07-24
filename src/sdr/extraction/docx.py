from __future__ import annotations

from pathlib import Path

import docx as python_docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from sdr.detection import DetectionResult

from .caption import pop_adjacent_caption
from .models import BlockType, ExtractedBlock, ExtractedDocument, QualityReport, failed_extraction


def extract(path: Path, detection: DetectionResult) -> ExtractedDocument:
    try:
        document = python_docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - corrupt docx must not crash extraction
        return failed_extraction(path, detection.doc_type, f"open_error: {exc}")

    try:
        body_children = list(document.element.body.iterchildren())
    except Exception as exc:  # noqa: BLE001
        return failed_extraction(path, detection.doc_type, f"body_read_error: {exc}")

    # document.paragraphs / document.tables are separate flat lists that lose
    # interleaving order; walking body children and mapping back to them is
    # what preserves the true reading order between prose and tables.
    paragraphs_by_elem = {p._p: p for p in document.paragraphs}
    tables_by_elem = {t._tbl: t for t in document.tables}

    blocks: list[ExtractedBlock] = []
    failures: list[str] = []
    tables_found = 0
    heading_stack: list[tuple[int, str]] = []
    order = 0

    for elem in body_children:
        if elem in paragraphs_by_elem:
            order = _handle_paragraph(paragraphs_by_elem[elem], blocks, heading_stack, order)
        elif elem in tables_by_elem:
            order, found = _handle_table(tables_by_elem[elem], blocks, heading_stack, order, failures)
            tables_found += found

    if not blocks:
        failures.append("no_content_recovered")

    chars_recovered = sum(len(b.text) for b in blocks)

    return ExtractedDocument(
        source_path=str(path),
        doc_type=detection.doc_type,
        blocks=tuple(blocks),
        quality=QualityReport(
            doc_type=detection.doc_type,
            pages_processed=1,
            tables_found=tables_found,
            chars_recovered=chars_recovered,
            failures=tuple(failures),
        ),
    )


def _handle_paragraph(
    para: Paragraph,
    blocks: list[ExtractedBlock],
    heading_stack: list[tuple[int, str]],
    order: int,
) -> int:
    text = para.text.strip()
    if not text:
        return order

    style_name = (para.style.name or "").lower() if para.style else ""
    if style_name.startswith("heading"):
        level = _heading_level(style_name)
        heading_stack[:] = [h for h in heading_stack if h[0] < level]
        heading_stack.append((level, text))
        block_type = BlockType.HEADING
    else:
        block_type = BlockType.PARAGRAPH

    blocks.append(
        ExtractedBlock(
            block_type=block_type,
            text=text,
            order=order,
            section_path=tuple(h[1] for h in heading_stack),
        )
    )
    return order + 1


def _handle_table(
    table: Table,
    blocks: list[ExtractedBlock],
    heading_stack: list[tuple[int, str]],
    order: int,
    failures: list[str],
) -> tuple[int, int]:
    rows = tuple(tuple(cell.text.strip() for cell in row.cells) for row in table.rows)
    if not any(any(cell for cell in row) for row in rows):
        failures.append(f"empty_table_at_order_{order}")
        return order, 0

    caption = pop_adjacent_caption(blocks)
    blocks.append(
        ExtractedBlock(
            block_type=BlockType.TABLE,
            text="\n".join(" | ".join(row) for row in rows),
            order=order,
            section_path=tuple(h[1] for h in heading_stack),
            table_rows=rows,
            caption=caption,
        )
    )
    return order + 1, 1


def _heading_level(style_name: str) -> int:
    digits = "".join(ch for ch in style_name if ch.isdigit())
    return int(digits) if digits else 1
