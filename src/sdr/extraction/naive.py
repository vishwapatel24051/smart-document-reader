from __future__ import annotations

from pathlib import Path

import docx as python_docx
import fitz  # PyMuPDF
from bs4 import BeautifulSoup

from sdr.detection import DocumentType, detect_file

from .models import BlockType, ExtractedBlock, ExtractedDocument, QualityReport, failed_extraction


def naive_extract(path: Path) -> ExtractedDocument:
    """A deliberately naive baseline: whatever a generic text loader would
    produce for each format - one flat text blob per document, no
    reading-order reconstruction, no table structure, no headings, no OCR
    fallback.

    This exists to be measured against extract() (Phase 2) by Phase 7's
    evaluation harness, not as a recommended extraction path - it's the
    exact failure mode this project's README describes as the problem.
    """
    detection = detect_file(path)

    try:
        if detection.doc_type in (DocumentType.PDF_TEXT, DocumentType.PDF_SCANNED):
            text = _naive_pdf_text(path)
        elif detection.doc_type == DocumentType.DOCX:
            text = _naive_docx_text(path)
        elif detection.doc_type == DocumentType.HTML:
            text = _naive_html_text(path)
        elif detection.doc_type == DocumentType.PLAIN_TEXT:
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            return failed_extraction(
                path,
                detection.doc_type,
                f"unsupported_or_undetected_format: confidence={detection.confidence:.2f}, "
                f"signals={detection.signals}",
            )
    except Exception as exc:  # noqa: BLE001 - a naive loader must not crash either
        return failed_extraction(path, detection.doc_type, f"naive_extract_failed: {exc}")

    text = text.strip()
    if not text:
        return ExtractedDocument(
            source_path=str(path),
            doc_type=detection.doc_type,
            blocks=(),
            quality=QualityReport(doc_type=detection.doc_type, failures=("no_content_recovered",)),
        )

    block = ExtractedBlock(block_type=BlockType.PARAGRAPH, text=text, order=0)
    return ExtractedDocument(
        source_path=str(path),
        doc_type=detection.doc_type,
        blocks=(block,),
        quality=QualityReport(doc_type=detection.doc_type, pages_processed=1, chars_recovered=len(text)),
    )


def _naive_pdf_text(path: Path) -> str:
    doc = fitz.open(str(path))
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def _naive_docx_text(path: Path) -> str:
    document = python_docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    # Tables aren't skipped, but they aren't kept structured either -
    # cells get flattened into a line of space-joined text, same as most
    # generic text loaders: the failure mode is "tables collapse into
    # unordered numbers," not "tables vanish."
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _naive_html_text(path: Path) -> str:
    soup = BeautifulSoup(path.read_bytes(), "lxml")
    return soup.get_text(separator=" ")
