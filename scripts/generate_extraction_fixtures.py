"""Generates the committed fixtures used by tests/test_extraction.py.

Run with: python scripts/generate_extraction_fixtures.py
Outputs are committed to the repo; re-run only if the fixture set needs to
change.
"""

from __future__ import annotations

from pathlib import Path

import docx
import fitz  # PyMuPDF

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "extraction"


def _write_bytes(name: str, data: bytes) -> None:
    path = FIXTURES_DIR / name
    path.write_bytes(data)
    print(f"wrote {path} ({len(data)} bytes)")


def make_multi_column_pdf() -> bytes:
    """Two columns whose blocks interleave in y-position but must be read
    column-by-column, not top-to-bottom across the whole page."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)

    left_lines = [
        (50, 80, "Left column, item one."),
        (50, 180, "Left column, item two."),
        (50, 280, "Left column, item three."),
    ]
    right_lines = [
        (400, 90, "Right column, item one."),
        (400, 190, "Right column, item two."),
        (400, 290, "Right column, item three."),
    ]
    for x, y, text in left_lines + right_lines:
        page.insert_text((x, y), text, fontsize=10)

    data = doc.tobytes(deflate=True, garbage=4)
    doc.close()
    return data


def make_table_pdf() -> bytes:
    """A captioned, ruled table - the shape PyMuPDF's find_tables() detects
    reliably, plus a heading and closing paragraph for reading-order context."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)

    page.insert_text((72, 72), "Experiment Results", fontsize=16)  # heading (large font)
    page.insert_text((72, 100), "Table 1: Example results", fontsize=11)  # caption

    x0, y0 = 72, 120
    col_w = [150, 150]
    row_h = 25
    rows, cols = 3, 2
    x_edges = [x0, x0 + col_w[0], x0 + col_w[0] + col_w[1]]
    y_edges = [y0 + i * row_h for i in range(rows + 1)]

    for y in y_edges:
        page.draw_line((x_edges[0], y), (x_edges[-1], y))
    for x in x_edges:
        page.draw_line((x, y_edges[0]), (x, y_edges[-1]))

    cell_text = [["Trial", "Outcome"], ["1", "alpha"], ["2", "beta"]]
    for r in range(rows):
        for c in range(cols):
            page.insert_text((x_edges[c] + 5, y_edges[r] + 17), cell_text[r][c], fontsize=10)

    page.insert_text((72, 230), "The table above summarizes the two trials that were run.", fontsize=11)

    data = doc.tobytes(deflate=True, garbage=4)
    doc.close()
    return data


def make_structured_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("This document exercises the docx extractor with headings, prose, and a table.")
    document.add_heading("Results", level=2)
    document.add_paragraph("Table 1: Summary of results")

    table = document.add_table(rows=3, cols=2)
    values = [["Trial", "Outcome"], ["1", "alpha"], ["2", "beta"]]
    for row, values_row in zip(table.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value

    document.add_paragraph("The table above summarizes the findings.")

    tmp_path = FIXTURES_DIR / "_tmp_structured.docx"
    document.save(tmp_path)
    data = tmp_path.read_bytes()
    tmp_path.unlink()
    return data


def make_structured_html() -> bytes:
    return (
        b"<!DOCTYPE html>\n<html><head><title>Fixture</title></head><body>"
        b"<h1>Introduction</h1>"
        b"<p>This document exercises the html extractor with headings, prose, and a table.</p>"
        b"<h2>Results</h2>"
        b"<table><caption>Table 1: Summary of results</caption>"
        b"<tr><th>Trial</th><th>Outcome</th></tr>"
        b"<tr><td>1</td><td>alpha</td></tr>"
        b"<tr><td>2</td><td>beta</td></tr>"
        b"</table>"
        b"<p>The table above summarizes the findings.</p>"
        b"</body></html>"
    )


def make_scanned_text_pdf() -> bytes:
    """An image-only page (no text layer) whose image contains real,
    legible rendered text - unlike detection's sample_scanned.pdf fixture,
    which is a blank pixel and can never OCR to anything. Lets the OCR
    success path be exercised wherever tesseract is actually installed."""
    text_doc = fitz.open()
    text_page = text_doc.new_page(width=400, height=200)
    text_page.insert_text((20, 100), "OCR RECOVERY TEST", fontsize=24)
    png_bytes = text_page.get_pixmap(dpi=200).tobytes("png")
    text_doc.close()

    doc = fitz.open()
    page = doc.new_page(width=400, height=200)
    page.insert_image(page.rect, stream=png_bytes)
    data = doc.tobytes(deflate=True, garbage=4)
    doc.close()
    return data


def make_structured_txt() -> bytes:
    return (
        b"This document exercises the plain text extractor.\n\n"
        b"It has a few short paragraphs separated by blank lines.\n\n"
        b"This is the third and final paragraph."
    )


def main() -> None:
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)

    _write_bytes("multi_column.pdf", make_multi_column_pdf())
    _write_bytes("table.pdf", make_table_pdf())
    _write_bytes("scanned_text.pdf", make_scanned_text_pdf())
    _write_bytes("structured.docx", make_structured_docx())
    _write_bytes("structured.html", make_structured_html())
    _write_bytes("structured.txt", make_structured_txt())


if __name__ == "__main__":
    main()
