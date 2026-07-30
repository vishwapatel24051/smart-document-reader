"""Generates the DOCX/PDF documents in the Phase 7 evaluation corpus
(eval/corpus/). The .txt and .html documents in that directory are
hand-written directly and aren't touched by this script.

Run with: python scripts/generate_eval_corpus.py

See eval/README.md for the disclosure on how this corpus and the paired
question set were authored.
"""

from __future__ import annotations

from pathlib import Path

import docx
import fitz  # PyMuPDF

CORPUS_DIR = Path(__file__).resolve().parent.parent / "eval" / "corpus"


def _write_bytes(name: str, data: bytes) -> None:
    path = CORPUS_DIR / name
    path.write_bytes(data)
    print(f"wrote {path} ({len(data)} bytes)")


def make_company_quarterly_report() -> bytes:
    """A fictional company's report - the numbers are synthetic, invented
    for this fixture, not real financial data about any real company."""
    document = docx.Document()
    document.add_heading("Acme Fictional Corp: Quarterly Report", level=1)
    document.add_paragraph(
        "Acme Fictional Corp is an invented company used only to test table "
        "extraction and citation in this project's evaluation corpus. The "
        "figures below are synthetic and do not describe any real business."
    )

    document.add_heading("Revenue by Quarter", level=2)
    document.add_paragraph("Table 1: Quarterly revenue for fiscal year 2025 (synthetic data, in millions of dollars)")
    table = document.add_table(rows=5, cols=2)
    rows = [
        ("Quarter", "Revenue ($M)"),
        ("Q1", "1.2"),
        ("Q2", "1.5"),
        ("Q3", "1.4"),
        ("Q4", "1.8"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value

    document.add_paragraph(
        "Revenue grew each quarter except Q3, which dipped slightly before "
        "Q4 became the highest-revenue quarter of the year."
    )

    tmp_path = CORPUS_DIR / "_tmp_company_report.docx"
    document.save(tmp_path)
    data = tmp_path.read_bytes()
    tmp_path.unlink()
    return data


def make_human_body_systems() -> bytes:
    document = docx.Document()
    document.add_heading("Human Body Systems", level=1)
    document.add_paragraph(
        "The human body is organized into several major organ systems that "
        "work together to keep it alive and functioning."
    )

    document.add_heading("Circulatory System", level=2)
    document.add_paragraph(
        "The circulatory system is made up of the heart, blood vessels, and "
        "blood. Its main job is to transport oxygen, nutrients, and hormones "
        "to cells throughout the body and to carry away waste products."
    )

    document.add_heading("Respiratory System", level=2)
    document.add_paragraph(
        "The respiratory system is centered on the lungs. Its main job is "
        "gas exchange: bringing oxygen into the blood and removing carbon "
        "dioxide from it."
    )

    document.add_heading("Digestive System", level=2)
    document.add_paragraph(
        "The digestive system includes the stomach and intestines. Its main "
        "job is to break down food into nutrients the body can absorb and "
        "use for energy."
    )

    tmp_path = CORPUS_DIR / "_tmp_human_body.docx"
    document.save(tmp_path)
    data = tmp_path.read_bytes()
    tmp_path.unlink()
    return data


def make_mount_everest_pdf() -> bytes:
    """A genuine two-column layout, to also exercise Phase 2's reading-order
    logic within the eval corpus, not just tests/fixtures/."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((50, 60), "Mount Everest", fontsize=18)

    left_col = (
        "Mount Everest is Earth's highest mountain above sea level, "
        "located in the Mahalangur Himal sub-range of the Himalayas. "
        "Its official height, as recognized in a 2020 joint remeasurement "
        "by China and Nepal, is 8,849 meters (29,032 feet).\n\n"
        "The mountain sits on the border between Nepal and the Tibet "
        "Autonomous Region of China. It is known as Sagarmatha in Nepali "
        "and Qomolangma in Tibetan."
    )
    right_col = (
        "The first confirmed summit of Mount Everest was achieved on "
        "May 29, 1953, by Edmund Hillary of New Zealand and Tenzing "
        "Norgay, a Sherpa mountaineer from Nepal.\n\n"
        "Everest is part of the Himalayan mountain range, which formed "
        "from the collision of the Indian and Eurasian tectonic plates "
        "and continues to rise a few millimeters each year."
    )

    # Kept narrow and clear of the page's 35%-65% width band (214-398 on a
    # 612pt-wide page) on purpose: a column whose text reaches into that
    # band defeats extraction's gutter-detection heuristic, which then
    # falls back to a plain top-to-bottom sort - the same mistake made (and
    # fixed) once already in tests/fixtures/extraction/multi_column.pdf.
    rect_left = fitz.Rect(50, 100, 205, 700)
    rect_right = fitz.Rect(400, 100, 560, 700)
    page.insert_textbox(rect_left, left_col, fontsize=10)
    page.insert_textbox(rect_right, right_col, fontsize=10)

    data = doc.tobytes(deflate=True, garbage=4)
    doc.close()
    return data


def make_coffee_brewing_guide_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    y = 60
    page.insert_text((50, y), "A Short Guide to Coffee Brewing Methods", fontsize=16)
    y += 40

    sections = [
        (
            "Pour-Over",
            "Pour-over is a manual brewing method where hot water is poured "
            "over coffee grounds held in a paper filter. It typically uses a "
            "medium grind and produces a clean, light-bodied cup.",
        ),
        (
            "French Press",
            "French press is an immersion method: coarsely ground coffee "
            "steeps directly in hot water, then a metal mesh plunger "
            "separates the grounds from the liquid. It produces a fuller, "
            "heavier-bodied cup than pour-over.",
        ),
        (
            "Espresso",
            "Espresso forces hot water through finely ground, tightly packed "
            "coffee under high pressure. It produces a small, concentrated "
            "shot that is the base for drinks like lattes and cappuccinos.",
        ),
    ]
    for heading, body in sections:
        page.insert_text((50, y), heading, fontsize=13)
        y += 22
        page.insert_textbox(fitz.Rect(50, y, 560, y + 70), body, fontsize=11)
        y += 90

    data = doc.tobytes(deflate=True, garbage=4)
    doc.close()
    return data


def make_board_game_collection_pdf() -> bytes:
    """A personal, synthetic log - the play counts and ratings are
    invented for this fixture, not real survey or sales data."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((50, 60), "Board Game Collection Log (personal, synthetic data)", fontsize=14)
    page.insert_text(
        (50, 85),
        "Table 1: Times played and personal rating, as logged by one fictional collector",
        fontsize=10,
    )

    x0, y0 = 50, 110
    col_w = [160, 120, 120]
    row_h = 25
    rows = 5
    x_edges = [x0, x0 + col_w[0], x0 + col_w[0] + col_w[1], x0 + col_w[0] + col_w[1] + col_w[2]]
    y_edges = [y0 + i * row_h for i in range(rows + 1)]

    for y in y_edges:
        page.draw_line((x_edges[0], y), (x_edges[-1], y))
    for x in x_edges:
        page.draw_line((x, y_edges[0]), (x, y_edges[-1]))

    cell_text = [
        ["Game", "Times Played", "Rating (/10)"],
        ["Catan", "14", "8"],
        ["Ticket to Ride", "9", "7"],
        ["Pandemic", "6", "9"],
        ["Codenames", "21", "8"],
    ]
    for r in range(rows):
        for c in range(3):
            page.insert_text((x_edges[c] + 5, y_edges[r] + 17), cell_text[r][c], fontsize=10)

    page.insert_text((50, y_edges[-1] + 30), "Codenames has been played the most often; Pandemic has the highest rating.", fontsize=11)

    data = doc.tobytes(deflate=True, garbage=4)
    doc.close()
    return data


def main() -> None:
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    _write_bytes("company_quarterly_report.docx", make_company_quarterly_report())
    _write_bytes("human_body_systems.docx", make_human_body_systems())
    _write_bytes("mount_everest.pdf", make_mount_everest_pdf())
    _write_bytes("coffee_brewing_guide.pdf", make_coffee_brewing_guide_pdf())
    _write_bytes("board_game_collection.pdf", make_board_game_collection_pdf())


if __name__ == "__main__":
    main()
